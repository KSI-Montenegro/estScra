#!/usr/bin/env python3
"""Verify the exported listing data against property records and source fixtures."""
import csv
import json
from pathlib import Path
from openpyxl import load_workbook
import extract_listing_details as d
import collect_images as g

def main():
    records=json.loads(Path('listing_details/listings.json').read_text())
    inputs=g.make_properties('image_collection/input_urls.json')
    assert len(records)==len(inputs)==1797
    assert {r['property_key'] for r in records}=={r['key'] for r in inputs}
    total_images=0
    forbidden=('Kontakt Forma','Prijavi Oglašivača','adsbygoogle','Vaša poruka je poslana','Detaljan opis:')
    for r in records:
        folder=Path('property_images')/r['source']/r['listing_id']
        original=json.loads((folder/'listing.json').read_text())
        assert r['listing_url']==original['listing_url']
        assert r['workbook_references']==original['workbook_references']
        assert r['image_urls']==original['images']
        assert r==json.loads((folder/'details.json').read_text())
        assert (folder/'description.txt').read_text()==((r['description_original']+'\n') if r['description_original'] else '')
        assert all(word not in (r['description_original'] or '') for word in forbidden)
        assert r.get('phone_e164')!='+38269257004'
        if r.get('phone_e164'):assert d.phone(r['phone_e164'])==r['phone_e164']
        if not r['source_page_available']:assert not r['description_original']
        if r['description_original']:assert r['field_sources']['description_original'].startswith('page:')
        total_images+=len(r['image_urls'])
    assert total_images==14522
    with open('listing_details/listings.csv',encoding='utf-8-sig',newline='') as f:
        rows=list(csv.DictReader(f))
    assert len(rows)==len(records)
    for row,r in zip(rows,records):
        assert row['property_key']==r['property_key']
        assert row['description_original']==(r['description_original'] or '')
    wb=load_workbook('listing_details/GRADEN_listings_complete.xlsx',read_only=True,data_only=False)
    sheet=wb['Listings'];values=list(sheet.iter_rows(values_only=True));keys=list(values[0]);desc=keys.index('description_original')
    assert len(values)==1798
    for row,r in zip(values[1:],records):assert row[desc]==(r['description_original'] or None)
    # Numeric/locale regressions that would corrupt price and size imports.
    for raw,expected in [('160000.00€',160000),('€580.000',580000),('1.200,50 EUR',1200.5),('175 000 m2',175000),('Cijena na upit',None)]:
        assert d.numeric(raw)==expected,(raw,d.numeric(raw))
    bykey={r['property_key']:r for r in records}
    a=bykey['estitor_629345'];assert a['price_amount']==450 and a['area_m2']==45 and a['bathrooms']==1 and a['deal_type']=='rent'
    b=bykey['patuljak_82084'];assert b['price_amount']==160000 and b['area_m2']==40 and b['city']=='Tivat'
    c=bykey['realitica_3949789'];assert c['price_amount']==580000 and c['area_m2']==138 and c['bathrooms']==3
    land=bykey['patuljak_80090'];assert land['price_amount'] is None and land['price_on_request'] and land['lot_area_m2']==175000 and land['area_m2'] is None
    assert 'Kontaktiraj Oglašivača' not in c['description_original']
    # A second ad from an owner must not inherit the first ad's price/city/area.
    sample={'source':'patuljak','listing_id':'2','title':None,'source_page_available':False,'image_urls':[],
            'field_sources':{},'quality_flags':[],'workbook_conflicts':[],'listing_url':'https://patuljak.me/oglas/second--2'}
    meta={'workbook_references':[{'sheet':'Outreach list','row':2,'column':'second_ad_url'}]}
    d.enrich(sample,meta,{('Outreach list',2):{'title':'First property','city':'Budva','m2':50,'price_eur':120000,'phone_e164':'+38267123456'}},{})
    assert sample.get('price_amount') is None and sample.get('area_m2') is None and sample.get('city') is None
    assert sample['phone_e164']=='+38267123456'
    report={'input_mapping':'passed','all_1797_property_folders':'passed','all_14522_image_urls_preserved':'passed',
            'original_description_files_match_json_csv_xlsx':'passed','source_text_excludes_site_ui':'passed',
            'price_and_area_locale_cases':'passed','three_source_fixture_details':'passed','secondary_ad_does_not_inherit_primary_property_fields':'passed',
            'descriptions_fabricated':0,'raw_html_caches_included':False}
    g.atomic('listing_details/validation.json',json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
